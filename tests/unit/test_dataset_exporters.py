import json
import os
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import numpy as np
import openpyxl
import pytest

from mri_dataset_tool.application.services.summary import summarize_study
from mri_dataset_tool.domain.models.serie import MRISeries
from mri_dataset_tool.domain.models.study import MRIStudy
from mri_dataset_tool.infrastructure.dataset.exporters import (
    to_json,
    to_markdown,
    to_xml,
    write_json,
    write_markdown,
    write_xlsx,
    write_xml,
)
from mri_dataset_tool.infrastructure.dataset.report import build_dataset_summary

SECRETS = ("SECRET", "1.2.840.99", "19700101", "C:\\private", "/home/private")
WRITERS = [write_json, write_markdown, write_xml, write_xlsx]


def _series(volume, modality="T1w", findings=()) -> MRISeries:
    return MRISeries(
        modality=modality,
        volume=volume,
        metadata={
            "shape": tuple(volume.shape),
            "voxel_size": (1.0, 1.0, 2.0),
            "affine": np.diag([1.0, 1.0, 2.0, 1.0]),
            "orientation": ("R", "A", "S"),
            "validation": {"status": "ok", "findings": list(findings)},
        },
    )


def _summary(classes=("healthy", "tumor"), label="tumor", modality="T1w"):
    a = np.arange(8, dtype=np.float32).reshape(2, 2, 2) / 3
    b = np.full((2, 2, 2, 2), np.nan, dtype=np.float32)
    entries = [
        ("01", "1", summarize_study(MRIStudy(study_uid=None, series=[_series(a, modality), _series(b, "dwi")]))),
        ("02", None, summarize_study(MRIStudy(study_uid=None, series=[_series(a, modality, ["CONSTANT_VOLUME"])]))),
    ]
    return build_dataset_summary(entries, {("01", "1"): label}, classes=classes)


def _from_xml(element: ET.Element, key: str = "dataset_report"):
    """Inverse of to_xml for the test data (strings that are not JSON stay strings)."""
    children = list(element)
    if not children:
        if element.text is None:
            return {}
        try:
            return json.loads(element.text)
        except ValueError:
            return element.text
    if key in ("modalities", "classes"):
        return {child.get("name"): _from_xml(child, child.get("name")) for child in children}
    if all(child.tag == "item" for child in children):
        return [_from_xml(child, "item") for child in children]
    return {child.tag: _from_xml(child, child.tag) for child in children}


def _xlsx_float(value):
    """What openpyxl keeps of a float (``%.16g``)."""
    return float("%.16g" % value) if isinstance(value, float) else value


def _sheet_rows(path: Path) -> dict[str, list[tuple]]:
    workbook = openpyxl.load_workbook(path)
    return {sheet.title: list(sheet.iter_rows(values_only=True)) for sheet in workbook.worksheets}


def test_xml_is_well_formed_and_holds_the_whole_summary() -> None:
    summary = _summary()
    text = to_xml(summary)
    assert text.startswith('<?xml version="1.0" encoding="UTF-8"?>\n<dataset_report>')
    assert _from_xml(ET.fromstring(text)) == summary
    assert to_xml(_summary()) == text


def test_xml_escapes_markup_in_names() -> None:
    summary = _summary(classes=["a&b<c>"], label="a&b<c>", modality='"<T1w>&\'')
    root = ET.fromstring(to_xml(summary))
    assert _from_xml(root) == summary
    assert {m.get("name") for m in root.iter("modality")} == {'"<T1w>&\'', "dwi"}


def test_xml_of_an_empty_dataset() -> None:
    summary = build_dataset_summary([])
    assert _from_xml(ET.fromstring(to_xml(summary))) == summary


def test_xlsx_has_one_sheet_per_section_with_the_json_figures(tmp_path: Path) -> None:
    summary = _summary()
    rows = _sheet_rows(write_xlsx(summary, tmp_path / "report.xlsx"))
    assert list(rows) == ["Totals", "Modalities", "Distributions", "Intensity", "Labels", "Warnings"]
    assert rows["Totals"] == [
        ("item", "count"), ("participants", 2), ("sessions", 2), ("studies", 2), ("series", 3),
        ("unlabeled_sessions", 1), ("unlabeled_participants", 1),
    ]
    t1 = summary["modalities"]["T1w"]["intensity"]
    assert rows["Intensity"][0] == ("modality", "min", "max", "mean", "std", "finite_voxels", "non_finite_voxels")
    assert rows["Intensity"][1] == ("T1w", *map(_xlsx_float, t1.values()))
    assert rows["Intensity"][2] == ("dwi", None, None, None, None, 0, 16)
    assert ("T1w", "shape", "[2, 2, 2]", 2) in rows["Distributions"]
    assert rows["Labels"] == [
        ("class", "participants", "sessions", "series"), ("healthy", 0, 0, 0), ("tumor", 1, 1, 2),
    ]
    assert rows["Warnings"][1:] == [(w["code"], w["count"]) for w in summary["warnings"]]


def test_same_figures_in_json_markdown_xml_and_xlsx(tmp_path: Path) -> None:
    summary = _summary()
    from_json = json.loads(to_json(summary))
    from_xml = _from_xml(ET.fromstring(to_xml(summary)))
    markdown = to_markdown(summary)
    header, *rows = _sheet_rows(write_xlsx(summary, tmp_path / "r.xlsx"))["Intensity"]
    from_xlsx = {row[0]: dict(zip(header[1:], row[1:])) for row in rows}
    for name, stats in summary["modalities"].items():
        figures = stats["intensity"]
        assert from_json["modalities"][name]["intensity"] == figures
        assert from_xml["modalities"][name]["intensity"] == figures
        assert from_xlsx[name] == {key: _xlsx_float(value) for key, value in figures.items()}
        assert "| " + " | ".join(json.dumps(v) for v in figures.values()) + " |" in markdown


@pytest.mark.parametrize("text", ["=1+1", "+cmd", "-2+3", "@SUM(A1:A2)", "=HYPERLINK(\"http://x\")"])
def test_xlsx_stores_formula_like_text_as_strings(tmp_path: Path, text: str) -> None:
    path = write_xlsx(_summary(classes=[text], label=text), tmp_path / "r.xlsx")
    sheet = openpyxl.load_workbook(path)["Labels"]
    assert sheet["A2"].value == text and sheet["A2"].data_type == "s"
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        assert not any("vbaProject" in n or "externalLink" in n for n in names)
        assert not any(b"<f>" in archive.read(n) or b"<f " in archive.read(n) for n in names)


@pytest.mark.parametrize("bad", ["A\ud800", "A\x01", "A\x1f"])
@pytest.mark.parametrize("writer", [write_xml, write_xlsx])
def test_unholdable_text_has_fixed_message_and_leaves_nothing(tmp_path: Path, writer, bad: str) -> None:
    with pytest.raises(ValueError, match=r"^Report content is invalid\.$") as info:
        writer(_summary(classes=[bad], label=bad), tmp_path / "r.out")
    assert info.value.__cause__ is None
    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize("value", [float("nan"), float("inf"), float("-inf")])
@pytest.mark.parametrize("writer", WRITERS)
def test_non_finite_values_are_rejected_like_json(tmp_path: Path, writer, value: float) -> None:
    summary = _summary()
    summary["modalities"]["T1w"]["intensity"]["mean"] = value
    with pytest.raises(ValueError) as expected:
        to_json(summary)
    with pytest.raises(ValueError) as info:
        writer(summary, tmp_path / "r.out")
    assert str(info.value) == str(expected.value)
    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize("writer", WRITERS)
def test_writers_do_not_overwrite(tmp_path: Path, writer) -> None:
    path = writer(_summary(), tmp_path / "report.out")
    before = path.read_bytes()
    with pytest.raises(FileExistsError, match=r"^Report file already exists\.$"):
        writer(_summary(), path)
    assert path.read_bytes() == before
    assert [p.name for p in tmp_path.iterdir()] == ["report.out"]


def test_no_phi_or_user_name_in_any_format(tmp_path: Path) -> None:
    series = _series(np.ones((2, 2, 2), np.float32))
    series.series_uid = "1.2.840.99.1"
    series.files = [Path("C:\\private\\sub-SECRET_T1w.nii.gz")]
    series.metadata.update({
        "PatientName": "SECRET^NAME",
        "StudyDate": "19700101",
        "source": "/home/private/x.nii.gz",
        "provenance": {"path": "/home/private/x"},
        "sidecar": {"RepetitionTime": 2.0, "PatientName": "SECRET^NAME"},
    })
    study = MRIStudy(study_uid="1.2.840.99.3", series=[series], subject_id="SECRET", session_id="SECRET")
    summary = build_dataset_summary([("SECRET", "SECRET", summarize_study(study))], {("SECRET", "SECRET"): "tumor"})
    texts = [path.read_bytes() for path in (writer(summary, tmp_path / f"r{i}") for i, writer in enumerate(WRITERS))]
    with zipfile.ZipFile(tmp_path / "r3") as archive:
        texts += [archive.read(name) for name in archive.namelist()]
    user = os.environ.get("USERNAME") or os.environ.get("USER")
    for content in texts:
        for secret in (*SECRETS, str(tmp_path), *([user] if user else [])):
            assert secret.encode() not in content
